import PyPDF2
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from bs4 import BeautifulSoup
import openpyxl
import xlrd
import csv

class ValidateDocumentFormat:
    def __init__(self, document_path: str):
        self.document = document_path
        self.document_type = ['PDF', 'DOCX', 'TEXT', 'HTML', 'XLS' ,'XLSX', 'CSV']

    def validate(self) -> bool:
        for doc_type in self.document_type:
           if self._is_valid_document_type(doc_type):
               return True
        return False
    
    def get_document_format(self) -> str:
        for doc_type in self.document_type:
            if self._is_valid_document_type(doc_type):
                return doc_type
        return 'UNKNOWN'
    
    def _is_valid_document_type(self, document_type: str) -> bool:
        match document_type:
            case 'PDF':
                return self._is_valid_pdf()
            case 'DOCX':
                return self._is_valid_docx()
            case 'TEXT':
                return self._is_valid_text()
            case 'HTML':
                return self._is_valid_html()
            case 'XLS':
                return self._is_valid_xls()
            case 'XLSX':
                return self._is_valid_xlsx()
            case 'CSV':
                return self._is_valid_csv()
            case _:
                return False
                
    def _is_valid_pdf(self) -> bool:
        try:
            with open(self.document, 'rb') as file:
                pdf = PyPDF2.PdfReader(file)
                if len(pdf.pages) > 0:
                    return True
                else:
                    return False
        except PyPDF2.errors.PdfReadError as e:
            return False
    
    def _is_valid_docx(self) -> bool:
        try:
            document = Document(self.document)
            return True
        except PackageNotFoundError:
            return False
        except Exception:
            return False
    
    def _is_valid_text(self) -> bool:
        try:
            with open(self.document, 'r', encoding='utf-8') as file:
                text = file.read()
                if len(text) > 0:
                    return True
                else:
                    return False
        except UnicodeDecodeError:
            return False
        except Exception:
            return False
    
    def _is_valid_html(self) -> bool:
        try:
            with open(self.document, 'r', encoding='utf-8') as file:
                html = file.read()
                soup = BeautifulSoup(html, 'html.parser')
                if soup.find_all('html'):
                    return True
                else:
                    return False
        except UnicodeDecodeError:
            return False
        except Exception:
            return False
    
    def _is_valid_xls(self) -> bool:
        try:
            workbook = xlrd.open_workbook(self.document)
            if len(workbook.sheets()) > 0:
                return True
            else:
                return False
        except xlrd.XLRDError:
            return False
        except Exception:
            return False
    
    def _is_valid_xlsx(self) -> bool:
        try:
            workbook = openpyxl.load_workbook(self.document)
            if len(workbook.sheetnames) > 0:
                return True
            else:
                return False
        except openpyxl.utils.exceptions.InvalidFileException:
            return False
        except Exception:
            return False
    
    def _is_valid_csv(self) -> bool:
        try:
            with open(self.document, 'r', newline='') as file:
                reader = csv.reader(file)
                if next(reader):
                    return True
                else:
                    return False
        except csv.Error:
            return False
        except Exception:
            return False